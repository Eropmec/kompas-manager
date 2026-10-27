"""Библиотека типовых элементов.

Библиотека — папка (в облаке) с подпапками по типам изделий. Элемент — модель КОМПАСа
(.m3d/.a3d) с внешними переменными. Рядом с моделью лежит паспорт «<модель>.kmlib.json»:
тип элемента, шаблоны обозначения и наименования, таблица типоразмеров, переменные.

В проект элемент попадает только копией: значения пишутся в переменные копии,
поэтому чертёж по ней строится с нужными размерами, а оригинал не меняется.
"""
import json
import os
from dataclasses import asdict, dataclass, field

import re

from .naming import file_name, format_number, format_template

_BAD_NAME = re.compile(r'[<>:"/\\|?*]')

PASSPORT_SUFFIX = ".kmlib.json"
MODEL_EXTS = (".m3d", ".a3d")
SKIP_DIRS = {"_old", "_backup", ".git"}

KINDS = {
    "standard": "Стандартное изделие",
    "like_standard": "Деталь по типу стандартной",
    "purchased": "Покупное изделие",
    "part": "Деталь",
}
SPEC_SECTIONS = {
    "standard": "Стандартные изделия",
    "like_standard": "Детали",
    "purchased": "Прочие изделия",
    "part": "Детали",
}


@dataclass
class Variable:
    name: str
    value: object = None
    note: str = ""
    info: bool = False  # информационная (вычисляемая), не задаётся при вставке


@dataclass
class Passport:
    kind: str = "standard"
    designation_template: str = ""
    name_template: str = ""
    table: str = ""  # таблица типоразмеров (.xlsx), путь относительно модели
    variables: list = field(default_factory=list)  # [Variable] — снимок из модели

    def editable_variables(self):
        return [v for v in self.variables if not v.info]

    def defaults(self):
        return {v.name: v.value for v in self.variables}

    def to_dict(self):
        d = asdict(self)
        return d

    @classmethod
    def from_dict(cls, d):
        p = cls(
            kind=d.get("kind", "standard") if d.get("kind") in KINDS else "standard",
            designation_template=d.get("designation_template", ""),
            name_template=d.get("name_template", ""),
            table=d.get("table", ""),
        )
        p.variables = [Variable(**{k: v.get(k) for k in ("name", "value", "note", "info") if k in v})
                       for v in d.get("variables", []) if v.get("name")]
        return p


def passport_path(model):
    return model + PASSPORT_SUFFIX


def load_passport(model):
    try:
        with open(passport_path(model), encoding="utf-8") as f:
            return Passport.from_dict(json.load(f))
    except (OSError, ValueError, TypeError):
        return None


def save_passport(model, passport):
    with open(passport_path(model), "w", encoding="utf-8") as f:
        json.dump(passport.to_dict(), f, ensure_ascii=False, indent=2)


@dataclass
class Element:
    path: str
    passport: Passport = None

    @property
    def title(self):
        return os.path.splitext(os.path.basename(self.path))[0]

    def table_path(self):
        if self.passport and self.passport.table:
            return os.path.normpath(os.path.join(os.path.dirname(self.path), self.passport.table))
        return ""


@dataclass
class Folder:
    path: str
    folders: list = field(default_factory=list)
    elements: list = field(default_factory=list)

    @property
    def title(self):
        return os.path.basename(self.path)

    def count(self):
        return len(self.elements) + sum(f.count() for f in self.folders)


def scan_library(root):
    """Дерево папок и элементов библиотеки (без КОМПАСа)."""
    def walk(path):
        folder = Folder(path)
        try:
            entries = sorted(os.scandir(path), key=lambda e: e.name.lower())
        except OSError:
            return folder
        for e in entries:
            if e.is_dir():
                if e.name in SKIP_DIRS or e.name.startswith("."):
                    continue
                sub = walk(e.path)
                if sub.count():
                    folder.folders.append(sub)
            elif os.path.splitext(e.name)[1].lower() in MODEL_EXTS and not e.name.startswith(("~", "_tmp_")):
                folder.elements.append(Element(e.path, load_passport(e.path)))
        return folder

    return walk(os.path.abspath(root))


# ---------- таблицы типоразмеров ----------

def read_table(path):
    """Первая страница xlsx: первая строка — имена столбцов (имена переменных), дальше типоразмеры.
    Возвращает (заголовки, строки-словари). Пустые строки пропускаются."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = None
        for row in rows:
            if row and any(c not in (None, "") for c in row):
                header = [str(c).strip() if c is not None else "" for c in row]
                break
        if not header:
            return [], []
        data = []
        for row in rows:
            if not row or all(c in (None, "") for c in row):
                continue
            data.append({h: _cell(v) for h, v in zip(header, row) if h})
        return [h for h in header if h], data
    finally:
        wb.close()


def _cell(v):
    if isinstance(v, str):
        s = v.strip().replace(",", ".")
        try:
            return float(s) if s else None
        except ValueError:
            return v.strip()
    return v


# ---------- вставка ----------

@dataclass
class InsertPlan:
    source: str
    target: str
    values: dict  # значения внешних переменных для записи в копию: {имя: значение}
    designation: str
    name: str
    reuse: bool = False  # такой файл уже есть в проекте — копия не нужна
    problems: list = field(default_factory=list)


def resolve_names(element, values, current_designation="", current_name=""):
    """Обозначение и наименование копии по шаблонам паспорта. Без шаблона — как у модели."""
    p = element.passport or Passport()
    all_values = dict(p.defaults())
    all_values.update(values)
    problems = []
    des, name = current_designation, current_name
    if p.designation_template:
        des, miss = format_template(p.designation_template, all_values)
        if miss:
            problems.append("В шаблоне обозначения нет значений для: " + ", ".join(miss))
    if p.name_template:
        name, miss = format_template(p.name_template, all_values)
        if miss:
            problems.append("В шаблоне наименования нет значений для: " + ", ".join(miss))
    return des.strip(), name.strip(), problems


NUMBERED_KINDS = ("like_standard", "part")  # получают обозначение по схеме изделия


def copy_file_name(designation, name, ext, title, values):
    """Имя файла копии: «Обозначение - Наименование»; без обозначения — по наименованию
    (стандартные и покупные изделия); без того и другого — по элементу и значениям."""
    if designation and name:
        return file_name(designation, name, ext)
    if name:
        return _BAD_NAME.sub("_", name.strip()) + ext
    suffix = " ".join(f"{k}{format_number(v)}" for k, v in values.items())
    return f"{title} {suffix}".strip() + ext


def values_key(values):
    """Ключ типоразмера для сравнения: значения в каноническом виде."""
    return {k: format_number(v) for k, v in sorted(values.items())}


def plan_insert(element, values, target_dir, current_designation="", current_name="", project_designation=""):
    """project_designation — обозначение по схеме изделия для деталей (тип «по типу стандартной», «деталь»),
    если в паспорте нет шаблона обозначения."""
    p = element.passport or Passport()
    known = {v.name for v in p.editable_variables()}
    clean = {}
    problems = []
    for n, v in values.items():
        if known and n not in known:
            problems.append(f"Нет внешней переменной {n}")
            continue
        if v in (None, ""):
            continue
        try:
            clean[n] = float(str(v).replace(",", "."))
        except ValueError:
            clean[n] = str(v)  # выражение КОМПАСа
    des, name, more = resolve_names(element, clean, current_designation, current_name)
    problems += more
    if p.kind in NUMBERED_KINDS and not p.designation_template and project_designation:
        des = project_designation
    ext = os.path.splitext(element.path)[1]
    target = os.path.join(target_dir, copy_file_name(des, name, ext, element.title, clean))
    if not name:
        problems.append("Не задано наименование — файл назван по модели и значениям")
    plan = InsertPlan(element.path, target, clean, des, name, problems=problems)
    plan.reuse = os.path.exists(target)
    if os.path.normcase(os.path.abspath(target)) == os.path.normcase(os.path.abspath(element.path)):
        plan.problems.append("Копия совпадает с оригиналом в библиотеке — выберите другую папку")
    return plan


# ---------- вынос вхождений с изменёнными переменными ----------

@dataclass
class ExtractGroup:
    source: str  # файл компонента
    values: dict  # значения, заданные в сборке
    target: str
    designation: str
    name: str
    instances: list = field(default_factory=list)  # индексы компонентов в сборке
    reuse: bool = False


def diff_values(in_assembly, in_file, eps=1e-9):
    out = {}
    for n, v in in_assembly.items():
        f = in_file.get(n)
        if isinstance(v, (int, float)) and isinstance(f, (int, float)):
            if abs(v - f) > eps:
                out[n] = v
        elif v != f:
            out[n] = v
    return out


def plan_extract(components, file_values, file_props, target_dir, passports=None):
    """components: [(индекс, файл, значения в сборке)]; file_values: файл -> значения в файле;
    file_props: файл -> (обозначение, наименование). Одинаковые значения — один новый файл."""
    passports = passports or {}
    groups = {}
    for idx, path, in_asm in components:
        changed = diff_values(in_asm, file_values.get(path, {}))
        if not changed:
            continue
        key = (os.path.normcase(path), tuple(sorted((k, format_number(v)) for k, v in changed.items())))
        if key not in groups:
            des0, name0 = file_props.get(path, ("", ""))
            element = Element(path, passports.get(path))
            if element.passport and (element.passport.designation_template or element.passport.name_template):
                merged = dict(file_values.get(path, {}))
                merged.update(in_asm)
                des, name, _ = resolve_names(Element(path, Passport(
                    designation_template=element.passport.designation_template,
                    name_template=element.passport.name_template)), merged, des0, name0)
            else:
                des, name = des0, name0
            ext = os.path.splitext(path)[1]
            suffix = " ".join(f"{k}{format_number(v)}" for k, v in sorted(changed.items()))
            if des and name and (des, name) != (des0, name0):
                target = os.path.join(target_dir, file_name(des, name, ext))
            else:
                stem = os.path.splitext(os.path.basename(path))[0]
                target = os.path.join(target_dir, f"{stem} [{suffix}]{ext}")
            groups[key] = ExtractGroup(path, changed, target, des, name)
            groups[key].reuse = os.path.exists(target)
        groups[key].instances.append(idx)
    return list(groups.values())
